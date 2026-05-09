from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas


class ReportGenerator:
    def to_markdown(self, analysis: dict, output_path: Path) -> Path:
        lines = [
            f"# {analysis.get('case_title', 'Case Report')}",
            "",
            "## Executive Summary",
            analysis.get("report_sections", {}).get("executive_summary", analysis.get("summary", "")),
            "",
            "## Findings",
        ]
        for finding in analysis.get("findings", []):
            lines.extend(
                [
                    f"### {finding.get('title', '')}",
                    f"- Severity: {finding.get('severity', '')}",
                    f"- Confidence: {finding.get('confidence', 0)}",
                    f"- Evidence IDs: {', '.join(finding.get('evidence_ids', []))}",
                    f"- Description: {finding.get('description', '')}",
                    f"- Impact: {finding.get('impact', '')}",
                    f"- Recommendation: {finding.get('recommendation', '')}",
                    "",
                ]
            )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("\n".join(lines), encoding="utf-8")
        return output_path

    def to_docx(self, analysis: dict, output_path: Path) -> Path:
        doc = Document()
        doc.add_heading(analysis.get("case_title", "Case Report"), level=1)
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(analysis.get("report_sections", {}).get("executive_summary", analysis.get("summary", "")))

        doc.add_heading("Findings", level=2)
        for finding in analysis.get("findings", []):
            doc.add_heading(finding.get("title", "Untitled"), level=3)
            doc.add_paragraph(f"Severity: {finding.get('severity', '')}")
            doc.add_paragraph(f"Confidence: {finding.get('confidence', 0)}")
            doc.add_paragraph(f"Evidence IDs: {', '.join(finding.get('evidence_ids', []))}")
            doc.add_paragraph(f"Description: {finding.get('description', '')}")
            doc.add_paragraph(f"Impact: {finding.get('impact', '')}")
            doc.add_paragraph(f"Recommendation: {finding.get('recommendation', '')}")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return output_path

    def to_pdf(self, analysis: dict, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        c = canvas.Canvas(str(output_path), pagesize=A4)
        _, height = A4
        y = height - 2 * cm

        def write_line(text: str, gap: float = 0.65 * cm):
            nonlocal y
            if y < 2 * cm:
                c.showPage()
                y = height - 2 * cm
            c.drawString(2 * cm, y, text[:120])
            y -= gap

        write_line(analysis.get("case_title", "Case Report"), 0.9 * cm)
        write_line("Executive Summary", 0.8 * cm)
        write_line(analysis.get("report_sections", {}).get("executive_summary", analysis.get("summary", "")))
        write_line("Findings", 0.8 * cm)

        for finding in analysis.get("findings", []):
            write_line(f"- {finding.get('title', 'Untitled')} [{finding.get('severity', '')}]", 0.7 * cm)
            write_line(f"  Confidence: {finding.get('confidence', 0)}")
            write_line(f"  Evidence IDs: {', '.join(finding.get('evidence_ids', []))}")
            write_line(f"  Description: {finding.get('description', '')}")

        c.save()
        return output_path
